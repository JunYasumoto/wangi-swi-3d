# -*- coding: utf-8 -*-
"""PowerPoint deck: 'Building a FloPy/MODFLOW environment with Claude Code'.
Environment-construction-centric; the Wangi SWI calibration is the closing showcase.
"""
import pathlib
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from PIL import Image

W = pathlib.Path(__file__).resolve().parent
FIG = W / "figs"

ACCENT = RGBColor(0xC1, 0x5F, 0x3C)
DARK   = RGBColor(0x2B, 0x2B, 0x2B)
GREY   = RGBColor(0x66, 0x66, 0x66)
LIGHT  = RGBColor(0xF2, 0xEE, 0xE9)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
CODEBG = RGBColor(0x1E, 0x1E, 0x1E)
CODEFG = RGBColor(0xEA, 0xEA, 0xEA)
CODECM = RGBColor(0x7E, 0xC6, 0x99)   # comment green
JP = "Yu Gothic UI"
MONO = "Consolas"

prs = Presentation()
prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def _set(run, size, color=DARK, bold=False, font=JP):
    run.font.size = Pt(size); run.font.bold = bold
    run.font.color.rgb = color; run.font.name = font


def box(slide, l, t, w, h, fill=None, line=None):
    sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    sp.shadow.inherit = False
    if fill is None: sp.fill.background()
    else: sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None: sp.line.fill.background()
    else: sp.line.color.rgb = line
    return sp


def text(slide, l, t, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space=6):
    tb = slide.shapes.add_textbox(l, t, w, h); tf = tb.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    for i, item in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(space)
        txt, kw = item; kw = dict(kw)
        p.level = kw.pop("level", 0)
        r = p.add_run(); r.text = txt; _set(r, **kw)
    return tb


def add_image_fit(slide, path, l, t, w, h):
    im = Image.open(path); ar = im.size[0] / im.size[1]
    if w / h > ar: nh = h; nw = int(h * ar)
    else: nw = w; nh = int(w / ar)
    slide.shapes.add_picture(str(path), l + (w - nw) // 2, t + (h - nh) // 2, nw, nh)


def header(slide, title, kicker=None):
    box(slide, 0, 0, SW, Inches(1.15), fill=DARK)
    box(slide, 0, Inches(1.15), SW, Pt(4), fill=ACCENT)
    text(slide, Inches(0.5), Inches(0.12), SW - Inches(1), Inches(0.9),
         [(title, dict(size=25, color=WHITE, bold=True))], anchor=MSO_ANCHOR.MIDDLE)
    if kicker:
        text(slide, Inches(0.55), Inches(0.14), Inches(5), Inches(0.3),
             [(kicker, dict(size=11, color=ACCENT, bold=True))])


def bullets_runs(bullets, big=20, small=16):
    runs = []
    for b in bullets:
        txt, lvl = (b if isinstance(b, tuple) else (b, 0))
        pre = "■ " if lvl == 0 else "– "
        runs.append((pre + txt, dict(size=big if lvl == 0 else small,
                     color=DARK if lvl == 0 else GREY, bold=(lvl == 0), level=lvl)))
    return runs


def content(title, bullets, kicker=None):
    s = prs.slides.add_slide(BLANK); header(s, title, kicker)
    text(s, Inches(0.7), Inches(1.5), SW - Inches(1.4), SH - Inches(2.0),
         bullets_runs(bullets), space=10)
    return s


def content_img(title, bullets, img, kicker=None, img_w=5.4):
    s = prs.slides.add_slide(BLANK); header(s, title, kicker)
    text(s, Inches(0.6), Inches(1.5), Inches(13.333 - img_w - 1.1), SH - Inches(2.0),
         bullets_runs(bullets, 18, 14), space=9)
    if (FIG / img).exists():
        add_image_fit(s, FIG / img, Inches(13.333 - img_w - 0.4), Inches(1.45),
                      Inches(img_w), SH - Inches(1.9))
    return s


def code_box(slide, l, t, w, h, lines):
    box(slide, l, t, w, h, fill=CODEBG)
    box(slide, l, t, Pt(6), h, fill=ACCENT)
    tb = slide.shapes.add_textbox(l + Inches(0.25), t + Inches(0.12),
                                  w - Inches(0.4), h - Inches(0.24))
    tf = tb.text_frame; tf.word_wrap = True
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(3)
        r = p.add_run(); r.text = ln if ln else " "
        col = CODECM if ln.strip().startswith("#") else CODEFG
        _set(r, 13, col, font=MONO)


def content_code(title, intro, code_lines, note=None, kicker=None, code_h=4.0):
    s = prs.slides.add_slide(BLANK); header(s, title, kicker)
    text(s, Inches(0.7), Inches(1.4), SW - Inches(1.4), Inches(1.0),
         bullets_runs(intro, 18, 15), space=6)
    ct = Inches(2.45)
    code_box(s, Inches(0.7), ct, SW - Inches(1.4), Inches(code_h), code_lines)
    if note:
        text(s, Inches(0.7), ct + Inches(code_h) + Inches(0.15), SW - Inches(1.4),
             Inches(0.8), [(note, dict(size=14, color=GREY))], space=4)
    return s


# ============================================================ 1 TITLE
s = prs.slides.add_slide(BLANK)
box(s, 0, 0, SW, SH, fill=DARK); box(s, 0, Inches(4.55), SW, Pt(5), fill=ACCENT)
text(s, Inches(0.9), Inches(1.9), SW - Inches(1.8), Inches(2.5), [
    ("Claude Code で始める", dict(size=34, color=WHITE, bold=True)),
    ("FloPy / MODFLOW 環境構築", dict(size=40, color=ACCENT, bold=True)),
], space=12)
text(s, Inches(0.95), Inches(4.8), SW - Inches(1.8), Inches(1.5), [
    ("AI エージェントと進める、地下水モデリングのセットアップ", dict(size=17, color=LIGHT)),
    ("conda 仮想環境 + flopy + MODFLOW 6 実行体 / 実例: Wangi 3D 海水侵入モデル",
     dict(size=13, color=GREY)),
], space=8)

# ============================================================ 2 主題
content("このトークの主題", [
    "ゼロから FloPy/MODFLOW 環境を立ち上げ、解析を回すまでを AI エージェントと",
    ("焦点は『環境構築』— 仮想環境・flopy・MODFLOW 実行体・動作確認", 1),
    "コマンド実行・バージョン整合・実行体入手・デバッグを Claude Code が代行",
    ("人は方針と判断に集中、AI が手を動かす", 1),
    "最後に実例の解析成果を少しだけ（Wangi 3D 海水侵入モデル）",
], kicker="INTRODUCTION")

# ============================================================ 3 Claude Code とは
content("Claude Code とは", [
    "ターミナルで動く AI コーディングエージェント (Anthropic)",
    ("コマンド実行 (PowerShell / bash)、ファイル編集、結果の確認・要約", 1),
    ("長時間ジョブのバックグラウンド実行、セッションをまたぐ永続メモリ", 1),
    "「環境構築のコマンドを、対話しながら実際に実行してくれる」",
    ("conda / pip / get-modflow を打ち、出力を読んで次の手を提案", 1),
    ("つまずけばエラーを解析して直す", 1),
], kicker="WHAT IS IT")

# ============================================================ 4 スタック図
content_img("FloPy / MODFLOW スタックの全体像", [
    "必要なものは大きく4層:",
    ("Python (conda 仮想環境)", 1),
    ("flopy = MODFLOW を Python で組む橋渡し", 1),
    ("MODFLOW 6 の実行体 (mf6.exe)", 1),
    ("解析 numpy/pandas/… と PEST++/pyemu", 1),
    "Claude Code が各層のセットアップ・実行・修正を代行",
], "flopy_stack.png", kicker="THE STACK", img_w=6.4)

# ============================================================ 5 なぜ大変か
content("環境構築が地味に大変な理由", [
    "パッケージのバージョン整合 (numpy / flopy / scipy …)",
    "MODFLOW は『実行体 (コンパイル済バイナリ)』が別途必要",
    ("Python ライブラリだけでは動かない", 1),
    "OS 差・特殊ビルド・絶対パス",
    ("例: 海水侵入(SWI)+潮汐は公式リリースに無い特殊ビルドが要る", 1),
    "→ この試行錯誤を AI に任せると速い",
], kicker="WHY IT'S FIDDLY")

# ============================================================ 6 Step1 conda
content_code("Step 1 — 仮想環境を作る (conda)", [
    "プロジェクト専用の隔離環境を作る → 既存環境を壊さない",
], [
    "# 仮想環境を作成して有効化",
    "conda create -n flopy_env python=3.11 -y",
    "conda activate flopy_env",
], note="Claude Code がコマンドを実行し、作成された環境を env list で確認。",
    kicker="STEP 1 / 4", code_h=2.2)

# ============================================================ 7 Step2 flopy
content_code("Step 2 — flopy と依存パッケージ", [
    "flopy 本体と、解析・キャリブレーションに使う一式を導入",
], [
    "# flopy + 解析 + PEST 連携ライブラリ",
    "pip install flopy matplotlib pandas scipy pyemu",
    "",
    "# 導入確認",
    'python -c "import flopy; print(flopy.__version__)"   # -> 3.9.2',
], note="バージョンを即確認。整合しなければ Claude Code が固定バージョンを提案。",
    kicker="STEP 2 / 4", code_h=3.0)

# ============================================================ 8 Step3 MODFLOW exe
content_code("Step 3 — MODFLOW 実行体を入手", [
    "flopy 同梱の get-modflow で USGS 公式バイナリを自動取得",
], [
    "# 公式 MODFLOW 実行体 (mf6.exe ほか) を ./bin にDL",
    "get-modflow ./bin",
    "",
    "# 本プロジェクトは SWI + 潮汐の特殊ビルドが必要",
    "#   → 公式 get-modflow では入らない",
    "#   → 自前/配布ビルドを配置:",
    "#     wangi_swi/mf6.6.0-SWI-Tides/mf6.exe",
], note="標準は get-modflow で完結。特殊機能だけ Claude Code が所在を探索・特定。",
    kicker="STEP 3 / 4", code_h=3.6)

# ============================================================ 9 Step4 smoke
content_code("Step 4 — 動作確認 (スモークテスト)", [
    "最小モデルを flopy で組んで実行 → success なら環境完成",
], [
    "import flopy",
    'sim = flopy.mf6.MFSimulation(sim_ws="smoke", exe_name="mf6")',
    "flopy.mf6.ModflowTdis(sim); flopy.mf6.ModflowIms(sim)",
    'gwf = flopy.mf6.ModflowGwf(sim, modelname="m")',
    "flopy.mf6.ModflowGwfdis(gwf, nlay=1, nrow=10, ncol=10)",
    "flopy.mf6.ModflowGwfic(gwf); flopy.mf6.ModflowGwfnpf(gwf)",
    "flopy.mf6.ModflowGwfchd(gwf,",
    "    stress_period_data=[[(0,0,0),1.0],[(0,9,9),0.0]])",
    'flopy.mf6.ModflowGwfoc(gwf, head_filerecord="m.hds",',
    '    saverecord=[("HEAD","ALL")])',
    "sim.write_simulation(); ok, _ = sim.run_simulation()",
    'print("success:", ok)    # -> True なら環境OK',
], kicker="STEP 4 / 4", code_h=4.6)

# ============================================================ 10 環境差異
content("実プロジェクトでの『環境差異』対応", [
    "本事例: 2 台の PC を Dropbox 同期 → 絶対パスがずれて動かない",
    ("python / mf6 / pestpp の場所がマシンで違う", 1),
    "Claude Code が自動対応:",
    ("ホスト名を検出し、制御ファイルのパスを書き換え", 1),
    ("SWI 対応の特殊 mf6 ビルドを探索して特定", 1),
    ("別マシンの過去ログとの差分から原因を即同定", 1),
    "→ 「なぜ動かないか」を人手で追う時間を圧縮",
], kicker="REAL WORLD")

# ============================================================ 11 再現性
content("壊さず・再現する工夫", [
    "環境を固定して共有・再現:",
    ("conda env export > environment.yml / pip freeze > requirements.txt", 1),
    "Claude Code の永続メモリに『環境メモ』を保存",
    ("python/mf6/pestpp のパス、特殊ビルドの所在を記録", 1),
    ("次セッションでも即座に環境差異へ対応できる", 1),
    "→ 一度きりのセットアップを、再現可能な資産に",
], kicker="REPRODUCIBILITY")

# ============================================================ 12 構図
content("ここまでの構図", [
    "人は『方針・判断』、AI は『手を動かす』",
    ("create → install → get-modflow → verify → fix を Claude Code が実行", 1),
    ("各ステップの出力を読み、次の一手を提案", 1),
    "環境構築の試行錯誤 (普段いちばん面倒な所) を圧縮",
    ("解析本体に早く入れる", 1),
], kicker="THE DIVISION OF LABOR")

# ============================================================ 13 showcase 1
content_img("そして解析へ — ショーケース①", [
    "環境ができたら flopy でモデル構築 → PEST++ でキャリブレーション",
    ("Wangi 島 3D 海水侵入モデル (MODFLOW 6 + SWI)", 1),
    ("観測井 20 本、水理状態で分類して同化集合を定義", 1),
    "regional / perched / boundary-controlled",
], "paper_well_classification.png", kicker="SHOWCASE", img_w=5.6)

# ============================================================ 14 showcase 2
content_img("ショーケース② — 主結果", [
    "宙水を除いた 14 井戸・3-zone K2 キャリブレーション",
    ("RMSE 1.75 m、非現実的な涵養を使わない", 1),
    "regional 井戸は 1:1 付近に整合",
    "環境構築 → 解析 → 可視化 → 文書化まで一気通貫",
    ("この図やスライドも Claude Code が生成", 1),
], "paper_obs_vs_sim.png", kicker="SHOWCASE", img_w=5.6)

# ============================================================ 15 効きどころ
content("Claude Code が環境構築で効いた点", [
    "conda / pip / get-modflow のコマンド実行と結果確認を代行",
    "MODFLOW 実行体の入手・特殊ビルドの探索",
    "環境差異 (パス・OS) の自動検出・修正",
    "動作確認 (スモークテスト) の自動化",
    "永続メモリ + env ファイルで再現性を担保",
    "セットアップから解析・可視化まで対話で一貫",
], kicker="WHY IT WORKED")

# ============================================================ 16 まとめ
s = prs.slides.add_slide(BLANK)
box(s, 0, 0, SW, SH, fill=DARK); box(s, 0, Inches(1.0), SW, Pt(4), fill=ACCENT)
text(s, Inches(0.8), Inches(0.3), SW - Inches(1.6), Inches(0.8),
     [("まとめ", dict(size=30, color=WHITE, bold=True))])
text(s, Inches(0.9), Inches(1.5), SW - Inches(1.8), Inches(3.0), [
    ("FloPy/MODFLOW 環境構築は『手順は単純だが詰まりやすい』", dict(size=22, color=WHITE, bold=True)),
    ("仮想環境作成 → flopy 導入 → 実行体入手 → 動作確認 を", dict(size=18, color=LIGHT)),
    ("  Claude Code がコマンド実行・検証・修正まで代行", dict(size=18, color=LIGHT)),
    ("特殊ビルドや環境差異といった『現場の壁』も AI が探索して解決", dict(size=18, color=LIGHT)),
], space=12)
box(s, Inches(0.9), Inches(5.1), SW - Inches(1.8), Inches(1.5), fill=RGBColor(0x3A, 0x3A, 0x3A))
box(s, Inches(0.9), Inches(5.1), Pt(6), Inches(1.5), fill=ACCENT)
text(s, Inches(1.15), Inches(5.25), SW - Inches(2.2), Inches(1.2), [
    ("Key message", dict(size=12, color=ACCENT, bold=True)),
    ("環境構築という一番面倒な入口を AI が引き受けることで、", dict(size=16, color=WHITE)),
    ("人は『どうモデル化し、何を信じるか』という本質に集中できる。", dict(size=16, color=WHITE)),
], space=6)

out = W / "ClaudeCode_FloPy_env_slides.pptx"
prs.save(str(out))
print("saved", out.name, "slides:", len(prs.slides._sldIdLst))
