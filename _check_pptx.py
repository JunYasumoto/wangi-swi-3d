import io,sys
sys.stdout=io.TextIOWrapper(sys.stdout.buffer,encoding="utf-8",errors="replace")
from pptx import Presentation
p=Presentation("ClaudeCode_FloPy_env_slides.pptx")
print("slides:",len(p.slides._sldIdLst))
for i,s in enumerate(p.slides,1):
    pics=len([sh for sh in s.shapes if sh.shape_type==13])
    head=""
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            head=sh.text_frame.text.strip().split("\n")[0][:38]; break
    print(f"  {i:2d}  pics={pics}  | {head}")
